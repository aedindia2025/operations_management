from __future__ import annotations

import datetime
import io
import uuid
from urllib.parse import quote

import openpyxl
from django.db import connection, transaction
from django.http import HttpResponse
from rest_framework import status
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from master.tenant import request_company_id


def _dictfetchall(cursor):
    columns = [col[0] for col in cursor.description]
    return [dict(zip(columns, row)) for row in cursor.fetchall()]


def _unique_rows(rows, key_builder):
    seen = set()
    unique = []
    for row in rows:
        key = key_builder(row)
        if key in seen:
            continue
        seen.add(key)
        unique.append(row)
    return unique


def _fmt_date(value, output_fmt: str = "%d-%m-%Y") -> str:
    if not value:
        return ""
    if isinstance(value, (datetime.date, datetime.datetime)):
        return value.strftime(output_fmt)
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%Y/%m/%d", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.datetime.strptime(text, fmt).strftime(output_fmt)
        except ValueError:
            continue
    return text


def _status_label(value) -> str:
    return {"0": "Pending", "1": "Mismatch", "2": "Verified"}.get(str(value or "0"), "Pending")


def _status_code_from_label(label: str) -> str:
    mapping = {
        "Pending": "0",
        "Mismatch": "1",
        "Mismatch/Rejected": "1",
        "Verified": "2",
    }
    return mapping.get(label, "0")


def _bool_file(value) -> bool:
    return bool(str(value or "").strip())


def _generate_unique_id() -> str:
    return uuid.uuid4().hex[:18]


def _installation_file_url(value) -> str:
    name = str(value or "").strip()
    return f"/api/master/installation/files/{quote(name)}/" if name else ""


def _dc_required_sql(install_alias: str, po_key_expr: str) -> str:
    return f"""
        CASE
            WHEN COALESCE(NULLIF(TRIM({install_alias}.dc_required), ''), '0') = '1'
                 OR EXISTS (
                     SELECT 1
                     FROM product_details_sub pds
                     WHERE pds.form_main_unique_id = {po_key_expr}
                       AND pds.is_delete = 0
                       AND COALESCE(NULLIF(TRIM(pds.document_required), ''), '0') = 'dc_required'
                     LIMIT 1
                 )
            THEN '1'
            ELSE '0'
        END
    """


def _parse_list_params(request):
    data = request.data if request.method.upper() == "POST" else request.query_params
    search_data = data.get("search", "")
    search = (
        data.get("search[value]")
        or (search_data.get("value", "") if isinstance(search_data, dict) else search_data)
        or ""
    )
    return {
        "draw": int(data.get("draw", 1) or 1),
        "start": int(data.get("start", 0) or 0),
        "length": int(data.get("length", 10) or 10),
        "search": str(search).strip(),
        "from_date": str(data.get("from_date", "") or "").strip(),
        "to_date": str(data.get("to_date", "") or "").strip(),
        "opt": str(data.get("opt", "") or "").strip(),
        "opt2": str(data.get("opt2", "") or data.get("opt", "") or "").strip(),
        "followed_by": str(data.get("followed_by", "") or data.get("team_mem", "") or data.get("snr_team_mem", "") or "").strip(),
        "team_mem": str(data.get("team_mem", "") or data.get("followed_by", "") or "").strip(),
        "verified_team_mem": str(data.get("verified_team_mem", "") or data.get("followed_by", "") or "").strip(),
        "screen_id_val": str(data.get("screen_id_val", "") or "").strip(),
        "user_type_unique_id": str(data.get("user_type_unique_id", "") or "").strip(),
        "tab": str(data.get("tab", "pending") or "pending").strip().lower(),
        "pending_type": str(data.get("pending_type", "ir") or "ir").strip().lower(),
        "company_id": "",
    }


def _date_column(opt: str) -> str:
    mapping = {
        "4": "v.po_date",
        "5": "v.invoice_date",
        "6": "v.dc_date",
        "po_date": "v.po_date",
        "invoice_date": "v.invoice_date",
        "dc_date": "v.dc_date",
        "PO Date": "v.po_date",
        "Invoice Date": "v.invoice_date",
        "DC Date": "v.dc_date",
    }
    return mapping.get(opt, "v.po_date")


def _paginate_rows(base_sql: str, select_sql: str, sql_params: list[object], start: int, length: int):
    with connection.cursor() as cur:
        cur.execute(f"SELECT COUNT(*) {base_sql}", sql_params)
        total = cur.fetchone()[0]

        query = select_sql
        params = list(sql_params)
        if length != -1:
            query += " LIMIT %s OFFSET %s"
            params.extend([length, start])
        elif start:
            query += " LIMIT 18446744073709551615 OFFSET %s"
            params.append(start)

        cur.execute(query, params)
        rows = _dictfetchall(cur)
    return total, rows


def _build_pending_query(params: dict):
    if params["pending_type"] == "ir":
        where = []
        sql_params: list[object] = []

        if params["from_date"] and params["to_date"]:
            if params["opt"] == "5":
                where.append("DATE(ids.invoice_date) BETWEEN %s AND %s")
            else:
                where.append("DATE(pf.po_date) BETWEEN %s AND %s")
            sql_params.extend([params["from_date"], params["to_date"]])

        team_member = params["team_mem"] or params["followed_by"]
        if team_member and team_member.lower() != "all":
            where.append(
                "COALESCE(team.staff_id, icm.team_mem, ids.team_mem, '') = %s"
            )
            sql_params.append(team_member)

        if params["search"]:
            like = f"%{params['search']}%"
            where.append(
                """(
                    COALESCE(pf.po_num, '') LIKE %s
                    OR COALESCE(ids.invoice_no, '') LIKE %s
                    OR COALESCE(ids.dc_number, '') LIKE %s
                    OR COALESCE(team.staff_name, icm.team_mem, ids.team_mem, '') LIKE %s
                    OR COALESCE(dep.department, dep_sub.ledger_name, icm.ledger_name, '') LIKE %s
                    OR COALESCE(cd.district_name, c.con_district, '') LIKE %s
                    OR COALESCE(cs.state_name, c.con_state_name, '') LIKE %s
                    OR COALESCE(c.con_address, '') LIKE %s
                )"""
            )
            sql_params.extend([like, like, like, like, like, like, like, like])

        where.extend(
            [
                "ids.document_verification_status = '0'",
                "ids.is_delete = 0",
                "idm.is_delete = 0",
                "idm.dc_delivery_status = '3'",
                "COALESCE(NULLIF(TRIM(ids.documents_type1), ''), '0') <> '0'",
            ]
        )
        if params.get("company_id"):
            where.append("pf.sess_company_id = %s")
            sql_params.append(params["company_id"])

        where_sql = " AND ".join(where)
        base_sql = f"""
            FROM installation_details_sublist ids
            LEFT JOIN invoice_creation_main icm
              ON icm.form_main_unique_id = ids.po_form_unique_id
             AND icm.consignee_unique_id = ids.consignee_unique_id
             AND icm.invoice_no = ids.invoice_no
             AND icm.dc_number = ids.dc_number
             AND icm.is_delete = 0
            LEFT JOIN installation_details idm
              ON idm.po_form_unique_id = ids.po_form_unique_id
             AND idm.consignee_unique_id = ids.consignee_unique_id
             AND idm.invoice_no = ids.invoice_no
             AND idm.dc_number = ids.dc_number
             AND idm.is_delete = 0
            LEFT JOIN dc_ir_doc_dispatch_details dc_ir_doc
              ON dc_ir_doc.po_form_unique_id = ids.po_form_unique_id
             AND dc_ir_doc.consignee_unique_id = ids.consignee_unique_id
             AND dc_ir_doc.invoice_no = ids.invoice_no
             AND dc_ir_doc.dc_number = ids.dc_number
             AND dc_ir_doc.is_delete = 0
            LEFT JOIN po_form pf
              ON ids.po_form_unique_id = pf.unique_id
             AND pf.is_delete = 0
            LEFT JOIN consignee_details_sub c
              ON c.unique_id = ids.consignee_unique_id
             AND c.is_delete = 0
            LEFT JOIN district_creation cd
              ON cd.unique_id = c.con_district
             AND cd.is_delete = 0
            LEFT JOIN state_creation cs
              ON cs.unique_id = c.con_state_name
             AND cs.is_delete = 0
            LEFT JOIN department_creation dep
              ON dep.unique_id = icm.ledger_name
             AND dep.is_delete = 0
            LEFT JOIN department_creation_sublist dep_sub
              ON dep_sub.unique_id = icm.ledger_name
             AND dep_sub.is_delete = 0
            LEFT JOIN user team
              ON team.staff_id = COALESCE(icm.team_mem, ids.team_mem)
             AND team.is_delete = 0
            WHERE {where_sql}
        """
        select_sql = f"""
            SELECT
                ids.po_form_unique_id AS form_main_unique_id,
                ids.consignee_unique_id,
                ids.unique_id AS ins_unique_id,
                pf.po_num,
                pf.po_date,
                COALESCE(team.staff_name, icm.team_mem, ids.team_mem, '') AS followed_by,
                COALESCE(dep.department, dep_sub.ledger_name, icm.ledger_name, '') AS ledger_name,
                COALESCE(cd.district_name, c.con_district, '') AS ledger_city,
                COALESCE(cs.state_name, c.con_state_name, '') AS ledger_state,
                ids.invoice_no,
                DATE_FORMAT(ids.invoice_date, '%%d-%%m-%%Y') AS invoice_date,
                ids.dc_number,
                COALESCE(icm.dc_date, '') AS dc_date,
                ids.documents_type AS dc_status,
                ids.dc_file AS dc_file,
                ids.documents_type1 AS ir_status,
                ids.ir_file AS ir_file,
                COALESCE((
                    SELECT DATE_FORMAT(ir_pod_date, '%%d-%%m-%%Y')
                    FROM dc_ir_doc_dispatch_details
                    WHERE invoice_no = ids.invoice_no
                      AND dc_number = ids.dc_number
                      AND is_delete = 0
                    LIMIT 1
                ), '') AS ir_pod_date,
                COALESCE(ids.document_verification_status, '0') AS document_verification_status
            {base_sql}
            GROUP BY ids.dc_number
            ORDER BY pf.po_date DESC
        """
        return base_sql, select_sql, sql_params

    where = ["COALESCE(ids.document_verification_status, '0') = '0'", "ids.is_delete = 0"]
    sql_params: list[object] = []
    if params.get("company_id"):
        where.append("p.sess_company_id = %s")
        sql_params.append(params["company_id"])

    if params["pending_type"] == "snr":
        where.extend(
            [
                "COALESCE(ids.documents_type1, '0') = '0'",
                "COALESCE(dd.dc_dispatch_mode, '') <> ''",
                "COALESCE(dd.snr_dispatch_mode, '') <> ''",
                "dd.is_delete = 0",
            ]
        )
    else:
        where.append("(COALESCE(ids.documents_type2, '0') = '0' OR COALESCE(ids.documents_type2, '') = '')")

    if params["from_date"] and params["to_date"]:
        date_column = {
            "po_date": "DATE(COALESCE(ids.po_date, p.po_date))",
            "invoice_date": "DATE(COALESCE(ids.invoice_date, icm.invoice_date))",
            "dc_date": "DATE(COALESCE(dl.dc_date, icm.dc_date))",
        }.get(_date_column(params["opt"]).split(".")[-1], "DATE(COALESCE(ids.po_date, p.po_date))")
        where.append(f"{date_column} BETWEEN %s AND %s")
        sql_params.extend([params["from_date"], params["to_date"]])

    if params["followed_by"] and params["followed_by"].lower() != "all":
        where.append("(COALESCE(u.staff_name, idm.team_mem, c.team_mem, '') = %s OR COALESCE(idm.team_mem, c.team_mem, '') = %s)")
        sql_params.extend([params["followed_by"], params["followed_by"]])

    if params["search"]:
        like = f"%{params['search']}%"
        where.append(
            """(
                COALESCE(ids.po_num, '') LIKE %s
                OR COALESCE(d.ledger_name, '') LIKE %s
                OR COALESCE(ids.invoice_no, '') LIKE %s
                OR COALESCE(ids.dc_number, '') LIKE %s
                OR COALESCE(u.staff_name, idm.team_mem, c.team_mem, '') LIKE %s
                OR COALESCE(cd.district_name, c.con_district, '') LIKE %s
                OR COALESCE(cs.state_name, c.con_state_name, '') LIKE %s
            )"""
        )
        sql_params.extend([like, like, like, like, like, like, like])

    base_sql = f"""
        FROM installation_details_sublist ids
        LEFT JOIN installation_details idm
          ON idm.po_form_unique_id = ids.po_form_unique_id
         AND idm.consignee_unique_id = ids.consignee_unique_id
         AND idm.invoice_no = ids.invoice_no
         AND idm.dc_number = ids.dc_number
         AND idm.is_delete = 0
        LEFT JOIN invoice_creation_main icm
          ON icm.form_main_unique_id = ids.po_form_unique_id
         AND icm.consignee_unique_id = ids.consignee_unique_id
         AND icm.invoice_no = ids.invoice_no
         AND icm.dc_number = ids.dc_number
         AND icm.is_delete = 0
        LEFT JOIN po_form p
          ON p.unique_id = ids.po_form_unique_id
         AND p.is_delete = 0
        LEFT JOIN consignee_details_sub c
          ON c.unique_id = ids.consignee_unique_id
         AND c.is_delete = 0
        LEFT JOIN district_creation cd
          ON cd.unique_id = c.con_district
         AND cd.is_delete = 0
        LEFT JOIN state_creation cs
          ON cs.unique_id = c.con_state_name
         AND cs.is_delete = 0
        LEFT JOIN department_creation d
          ON d.unique_id = p.department
         AND d.is_delete = 0
        LEFT JOIN dispatch_list dl
          ON dl.invoice_no = ids.invoice_no
         AND dl.dc_number = ids.dc_number
         AND dl.consignee_unique_id = ids.consignee_unique_id
         AND dl.is_delete = 0
        LEFT JOIN dc_ir_doc_dispatch_details dd
          ON dd.po_form_unique_id = ids.po_form_unique_id
         AND dd.consignee_unique_id = ids.consignee_unique_id
         AND dd.invoice_no = ids.invoice_no
         AND dd.dc_number = ids.dc_number
         AND dd.is_delete = 0
        LEFT JOIN `user` u
          ON u.staff_id = COALESCE(idm.team_mem, c.team_mem)
         AND u.is_delete = 0
        WHERE {' AND '.join(where)}
    """

    select_sql = f"""
        SELECT
            ids.po_form_unique_id AS form_main_unique_id,
            ids.consignee_unique_id,
            ids.unique_id AS ins_unique_id,
            COALESCE(ids.po_num, p.po_num, '') AS po_num,
            COALESCE(ids.po_date, p.po_date, '') AS po_date,
            COALESCE(u.staff_name, icm.team_mem, idm.team_mem, c.team_mem, '') AS followed_by,
            COALESCE(d.department, '') AS ledger_name,
            COALESCE(cd.district_name, c.con_district, '') AS ledger_city,
            COALESCE(cs.state_name, c.con_state_name, '') AS ledger_state,
            ids.invoice_no,
            COALESCE(ids.invoice_date, icm.invoice_date, '') AS invoice_date,
            ids.dc_number,
            COALESCE(dl.dc_date, icm.dc_date) AS dc_date,
            COALESCE(ids.documents_type, '0') AS dc_status,
            COALESCE(ids.documents_type1, '0') AS ir_status,
            COALESCE(ids.documents_type2, '0') AS snr_status,
            COALESCE(ids.dc_file, idm.dc_file, '') AS dc_file,
            COALESCE(ids.ir_file, idm.ir_file, '') AS ir_file,
            COALESCE(ids.snr_file, idm.snr_file, '') AS snr_file,
            COALESCE((
                SELECT DATE_FORMAT(snr_pod_date, '%%d-%%m-%%Y')
                FROM dc_ir_doc_dispatch_details
                WHERE invoice_no = ids.invoice_no
                  AND dc_number = ids.dc_number
                  AND is_delete = 0
                LIMIT 1
            ), '') AS snr_pod_date,
            COALESCE(ids.document_verification_status, '0') AS document_verification_status
        {base_sql}
        GROUP BY ids.dc_number
        ORDER BY ids.id DESC
    """
    return base_sql, select_sql, sql_params


def _build_saved_query(params: dict, statuses: list[str]):
    where = [f"sd.status_app IN ({','.join(['%s'] * len(statuses))})", "sd.is_delete = 0"]
    sql_params: list[object] = list(statuses)
    if params.get("company_id"):
        where.append("pf.sess_company_id = %s")
        sql_params.append(params["company_id"])

    is_verified_tab = statuses == ["2"]
    date_opt = params.get("opt2") if is_verified_tab else params.get("opt")

    if params["from_date"] and params["to_date"]:
        saved_date_column = {
            "po_date": "DATE(COALESCE(NULLIF(sd.po_date, ''), ids.po_date, pf.po_date))",
            "invoice_date": "DATE(COALESCE(NULLIF(sd.invoice_date, ''), ids.invoice_date, icm.invoice_date))",
            "dc_date": "DATE(COALESCE(dl.dc_date, icm.dc_date))",
        }.get(_date_column(date_opt).split(".")[-1], "DATE(COALESCE(NULLIF(sd.po_date, ''), ids.po_date, pf.po_date))")
        where.append(f"{saved_date_column} BETWEEN %s AND %s")
        sql_params.extend([params["from_date"], params["to_date"]])

    team_filter = params.get("verified_team_mem") if is_verified_tab else params.get("followed_by")
    if team_filter and team_filter.lower() != "all":
        where.append("(COALESCE(u.staff_name, icm.team_mem, ids.team_mem, '') = %s OR COALESCE(icm.team_mem, ids.team_mem, '') = %s)")
        sql_params.extend([team_filter, team_filter])

    if params["search"]:
        like = f"%{params['search']}%"
        where.append(
            """(
                COALESCE(sd.po_num, '') LIKE %s
                OR COALESCE(dep.department, dep_sub.ledger_name, icm.ledger_name, '') LIKE %s
                OR COALESCE(sd.invoice_no, '') LIKE %s
                OR COALESCE(sd.dc_number, '') LIKE %s
                OR COALESCE(u.staff_name, icm.team_mem, ids.team_mem, '') LIKE %s
                OR COALESCE(cd.district_name, c.con_district, '') LIKE %s
                OR COALESCE(cs.state_name, c.con_state_name, '') LIKE %s
                OR COALESCE(sd.reject_reason, '') LIKE %s
            )"""
        )
        sql_params.extend([like, like, like, like, like, like, like, like])

    base_sql = f"""
        FROM sign_doc_verification_detail sd
        LEFT JOIN installation_details_sublist ids
          ON ids.unique_id = sd.ins_unique_id
         AND ids.po_form_unique_id = sd.form_main_unique_id
         AND ids.consignee_unique_id = sd.con_unique_id
         AND ids.invoice_no = sd.invoice_no
         AND ids.dc_number = sd.dc_number
         AND ids.is_delete = 0
        LEFT JOIN invoice_creation_main icm
          ON icm.form_main_unique_id = sd.form_main_unique_id
         AND icm.consignee_unique_id = sd.con_unique_id
         AND icm.invoice_no = sd.invoice_no
         AND icm.dc_number = sd.dc_number
         AND icm.is_delete = 0
        LEFT JOIN po_form pf
          ON pf.unique_id = sd.form_main_unique_id
         AND pf.is_delete = 0
        LEFT JOIN consignee_details_sub c
          ON c.unique_id = sd.con_unique_id
         AND c.is_delete = 0
        LEFT JOIN district_creation cd
          ON cd.unique_id = c.con_district
         AND cd.is_delete = 0
        LEFT JOIN state_creation cs
          ON cs.unique_id = c.con_state_name
         AND cs.is_delete = 0
        LEFT JOIN department_creation dep
          ON dep.unique_id = icm.ledger_name
         AND dep.is_delete = 0
        LEFT JOIN department_creation_sublist dep_sub
          ON dep_sub.unique_id = icm.ledger_name
         AND dep_sub.is_delete = 0
        LEFT JOIN dispatch_list dl
          ON dl.invoice_no = sd.invoice_no
         AND dl.dc_number = sd.dc_number
         AND dl.consignee_unique_id = sd.con_unique_id
         AND dl.is_delete = 0
        LEFT JOIN user u
          ON u.staff_id = COALESCE(icm.team_mem, ids.team_mem)
         AND u.is_delete = 0
        WHERE {' AND '.join(where)}
    """
    select_sql = f"""
        SELECT
            sd.unique_id AS verification_unique_id,
            sd.form_main_unique_id,
            sd.con_unique_id AS consignee_unique_id,
            sd.ins_unique_id,
            COALESCE(sd.po_num, ids.po_num, pf.po_num, '') AS po_num,
            COALESCE(NULLIF(sd.po_date, ''), ids.po_date, DATE_FORMAT(pf.po_date, '%%Y-%%m-%%d')) AS po_date,
            COALESCE(u.staff_name, icm.team_mem, ids.team_mem, '') AS followed_by,
            COALESCE(dep.department, dep_sub.ledger_name, icm.ledger_name, '') AS ledger_name,
            COALESCE(cd.district_name, c.con_district, '') AS ledger_city,
            COALESCE(cs.state_name, c.con_state_name, '') AS ledger_state,
            sd.invoice_no,
            COALESCE(NULLIF(sd.invoice_date, ''), ids.invoice_date, DATE_FORMAT(icm.invoice_date, '%%Y-%%m-%%d')) AS invoice_date,
            sd.dc_number,
            COALESCE(DATE_FORMAT(dl.dc_date, '%%Y-%%m-%%d'), DATE_FORMAT(icm.dc_date, '%%Y-%%m-%%d'), '') AS dc_date,
            COALESCE(sd.dc_received_status, '') AS dc_status,
            COALESCE(sd.ir_status, '') AS ir_status,
            COALESCE(sd.snr_status, '') AS snr_status,
            COALESCE(ids.dc_file, '') AS dc_file,
            COALESCE(ids.ir_file, '') AS ir_file,
            COALESCE(ids.snr_file, '') AS snr_file,
            COALESCE(sd.status_app, '0') AS status_app,
            COALESCE(sd.reject_reason, '') AS reject_reason
        {base_sql}
        ORDER BY sd.created DESC, sd.id DESC
    """
    return base_sql, select_sql, sql_params


def _build_verified_query(params: dict):
    where = ["ids.document_verification_status = 2", "ids.is_delete = 0"]
    sql_params: list[object] = []
    if params.get("company_id"):
        where.append("pf.sess_company_id = %s")
        sql_params.append(params["company_id"])

    if params["from_date"] and params["to_date"]:
        verified_date_column = {
            "100": "DATE(COALESCE(ids.po_date, pf.po_date))",
            "102": "DATE(COALESCE(ids.invoice_date, icm.invoice_date))",
            "po_date": "DATE(COALESCE(ids.po_date, pf.po_date))",
            "invoice_date": "DATE(COALESCE(ids.invoice_date, icm.invoice_date))",
        }.get(str(params.get("opt2") or "").strip(), "DATE(COALESCE(ids.po_date, pf.po_date))")
        where.append(f"{verified_date_column} BETWEEN %s AND %s")
        sql_params.extend([params["from_date"], params["to_date"]])

    team_filter = str(params.get("verified_team_mem") or "").strip()
    if team_filter and team_filter.lower() != "all":
        where.append("(COALESCE(u.staff_name, ids.team_mem, '') = %s OR COALESCE(ids.team_mem, '') = %s)")
        sql_params.extend([team_filter, team_filter])

    if params["search"]:
        like = f"%{params['search']}%"
        where.append(
            """(
                COALESCE(ids.po_num, '') LIKE %s
                OR COALESCE(dep.department, dep_sub.ledger_name, icm.ledger_name, '') LIKE %s
                OR COALESCE(u.staff_name, ids.team_mem, '') LIKE %s
                OR COALESCE(ids.invoice_no, '') LIKE %s
                OR COALESCE(ids.dc_number, '') LIKE %s
                OR COALESCE(cd.district_name, c.con_district, '') LIKE %s
                OR COALESCE(cs.state_name, c.con_state_name, '') LIKE %s
                OR COALESCE(c.con_address, '') LIKE %s
            )"""
        )
        sql_params.extend([like, like, like, like, like, like, like, like])

    base_sql = f"""
        FROM installation_details_sublist ids
        LEFT JOIN sign_doc_verification_detail sd
          ON sd.ins_unique_id = ids.unique_id
         AND sd.form_main_unique_id = ids.po_form_unique_id
         AND sd.con_unique_id = ids.consignee_unique_id
         AND sd.invoice_no = ids.invoice_no
         AND sd.dc_number = ids.dc_number
         AND sd.is_delete = 0
        LEFT JOIN invoice_creation_main icm
          ON icm.form_main_unique_id = ids.po_form_unique_id
         AND icm.dc_number = ids.dc_number
         AND icm.is_delete = 0
        LEFT JOIN po_form pf
          ON pf.unique_id = ids.po_form_unique_id
         AND pf.is_delete = 0
        LEFT JOIN consignee_details_sub c
          ON c.unique_id = ids.consignee_unique_id
         AND c.is_delete = 0
        LEFT JOIN district_creation cd
          ON cd.unique_id = c.con_district
         AND cd.is_delete = 0
        LEFT JOIN state_creation cs
          ON cs.unique_id = c.con_state_name
         AND cs.is_delete = 0
        LEFT JOIN department_creation dep
          ON dep.unique_id = icm.ledger_name
         AND dep.is_delete = 0
        LEFT JOIN department_creation_sublist dep_sub
          ON dep_sub.unique_id = icm.ledger_name
         AND dep_sub.is_delete = 0
        LEFT JOIN user u
          ON u.staff_id = ids.team_mem
         AND u.is_delete = 0
        WHERE {' AND '.join(where)}
    """

    select_sql = f"""
        SELECT
            sd.unique_id AS verification_unique_id,
            ids.po_form_unique_id AS form_main_unique_id,
            ids.consignee_unique_id,
            ids.unique_id AS ins_unique_id,
            COALESCE(ids.po_num, pf.po_num, '') AS po_num,
            COALESCE(ids.po_date, DATE_FORMAT(pf.po_date, '%%Y-%%m-%%d'), '') AS po_date,
            COALESCE(u.staff_name, ids.team_mem, '') AS followed_by,
            COALESCE(dep.department, dep_sub.ledger_name, icm.ledger_name, '') AS ledger_name,
            COALESCE(cd.district_name, c.con_district, '') AS ledger_city,
            COALESCE(cs.state_name, c.con_state_name, '') AS ledger_state,
            ids.invoice_no,
            COALESCE(ids.invoice_date, DATE_FORMAT(icm.invoice_date, '%%Y-%%m-%%d'), '') AS invoice_date,
            ids.dc_number,
            COALESCE(DATE_FORMAT(icm.dc_date, '%%Y-%%m-%%d'), '') AS dc_date,
            COALESCE(ids.dc_file, '') AS dc_file,
            COALESCE(ids.ir_file, '') AS ir_file,
            COALESCE(ids.snr_file, '') AS snr_file,
            COALESCE(ids.document_verification_status, '2') AS status_app,
            '' AS reject_reason,
            COALESCE(sd.bill_no, '') AS bill_no
        {base_sql}
        ORDER BY ids.id DESC
    """
    return base_sql, select_sql, sql_params


def _get_pending_ir_rows(params: dict):
    dc_req_exists_sql = _dc_required_sql("dc_req", "ids.po_form_unique_id")
    dc_required_sql = _dc_required_sql("idm", "ids.po_form_unique_id")
    conditions = [
        "ids.is_delete = 0",
        "COALESCE(ids.document_verification_status, '0') = '0'",
        f"""
        (
            COALESCE(NULLIF(TRIM(ids.documents_type1), ''), '0') <> '0'
            OR EXISTS (
                SELECT 1
                FROM installation_details dc_req
                WHERE dc_req.po_form_unique_id = ids.po_form_unique_id
                  AND dc_req.consignee_unique_id = ids.consignee_unique_id
                  AND dc_req.invoice_no = ids.invoice_no
                  AND dc_req.dc_number = ids.dc_number
                  AND dc_req.is_delete = 0
                  AND {dc_req_exists_sql} = '1'
                  AND COALESCE(NULLIF(TRIM(ids.documents_type), ''), '0') <> '0'
            )
        )
        """,
        f"""
        EXISTS (
            SELECT 1
            FROM installation_details idm
            WHERE idm.po_form_unique_id = ids.po_form_unique_id
              AND idm.consignee_unique_id = ids.consignee_unique_id
              AND idm.invoice_no = ids.invoice_no
              AND idm.dc_number = ids.dc_number
              AND idm.is_delete = 0
              AND (
                  idm.dc_delivery_status = '3'
                  OR ({dc_required_sql} = '1' AND idm.dc_delivery_status = '5')
              )
        )
        """,
    ]
    sql_params: list[object] = []
    if params.get("company_id"):
        conditions.append(
            """
            EXISTS (
                SELECT 1
                FROM po_form pf_tenant
                WHERE pf_tenant.unique_id = ids.po_form_unique_id
                  AND pf_tenant.is_delete = 0
                  AND pf_tenant.sess_company_id = %s
            )
            """
        )
        sql_params.append(params["company_id"])

    if params["from_date"] and params["to_date"]:
        if params["opt"] == "5":
            conditions.append("DATE(ids.invoice_date) BETWEEN %s AND %s")
        else:
            conditions.append("DATE(ids.po_date) BETWEEN %s AND %s")
        sql_params.extend([params["from_date"], params["to_date"]])

    team_member = params["team_mem"] or params["followed_by"]
    if team_member and team_member.lower() != "all":
        conditions.append(
            """
            (
                COALESCE(NULLIF(TRIM(ids.team_mem), ''), '') = %s
                OR EXISTS (
                    SELECT 1
                    FROM user team
                    WHERE (
                        team.staff_id = COALESCE(NULLIF(TRIM(ids.team_mem), ''), '')
                        OR team.unique_id = COALESCE(NULLIF(TRIM(ids.team_mem), ''), '')
                        OR team.user_name = COALESCE(NULLIF(TRIM(ids.team_mem), ''), '')
                    )
                      AND team.is_delete = 0
                      AND team.staff_name = %s
                )
            )
            """
        )
        sql_params.extend([team_member, team_member])

    if params["search"]:
        like = f"%{params['search']}%"
        conditions.append(
            """
            (
                COALESCE(ids.po_num, '') LIKE %s
                OR COALESCE(ids.invoice_no, '') LIKE %s
                OR COALESCE(ids.dc_number, '') LIKE %s
                OR EXISTS (
                    SELECT 1
                    FROM user team
                    WHERE (
                        team.staff_id = COALESCE(NULLIF(TRIM(ids.team_mem), ''), '')
                        OR team.unique_id = COALESCE(NULLIF(TRIM(ids.team_mem), ''), '')
                        OR team.user_name = COALESCE(NULLIF(TRIM(ids.team_mem), ''), '')
                    )
                      AND team.is_delete = 0
                      AND team.staff_name LIKE %s
                )
                OR EXISTS (
                    SELECT 1
                    FROM invoice_creation_main icm
                    LEFT JOIN department_creation dep
                      ON dep.unique_id = icm.ledger_name
                     AND dep.is_delete = 0
                    LEFT JOIN department_creation_sublist dep_sub
                      ON dep_sub.unique_id = icm.ledger_name
                     AND dep_sub.is_delete = 0
                    LEFT JOIN po_form pf
                      ON pf.unique_id = icm.form_main_unique_id
                     AND pf.is_delete = 0
                    LEFT JOIN department_creation dep_pf
                      ON dep_pf.unique_id = pf.department
                     AND dep_pf.is_delete = 0
                    WHERE icm.form_main_unique_id = ids.po_form_unique_id
                      AND icm.consignee_unique_id = ids.consignee_unique_id
                      AND icm.invoice_no = ids.invoice_no
                      AND icm.dc_number = ids.dc_number
                      AND icm.is_delete = 0
                      AND COALESCE(dep_sub.ledger_name, dep.ledger_name, dep.department, dep_pf.department, icm.ledger_name, pf.department, '') LIKE %s
                )
                OR EXISTS (
                    SELECT 1
                    FROM consignee_details_sub c
                    WHERE c.unique_id = ids.consignee_unique_id
                      AND c.is_delete = 0
                      AND (
                          COALESCE(c.con_district, '') LIKE %s
                          OR COALESCE(c.con_state_name, '') LIKE %s
                          OR COALESCE(c.con_address, '') LIKE %s
                      )
                )
            )
            """
        )
        sql_params.extend([like, like, like, like, like, like, like, like])

    where_sql = " AND ".join(conditions)

    count_sql = f"""
        SELECT COUNT(*)
        FROM (
            SELECT 1
            FROM installation_details_sublist ids
            WHERE {where_sql}
            GROUP BY ids.po_form_unique_id, ids.consignee_unique_id, ids.invoice_no, ids.dc_number
        ) q
    """

    page_sql = f"""
        SELECT
            MAX(ids.id) AS latest_id,
            ids.po_form_unique_id,
            ids.consignee_unique_id,
            ids.invoice_no,
            ids.dc_number
        FROM installation_details_sublist ids
        WHERE {where_sql}
        GROUP BY ids.po_form_unique_id, ids.consignee_unique_id, ids.invoice_no, ids.dc_number
        ORDER BY MAX(ids.id) DESC
    """

    with connection.cursor() as cur:
        cur.execute(count_sql, sql_params)
        total = cur.fetchone()[0] or 0

        page_query = page_sql
        page_params = list(sql_params)
        if params["length"] != -1:
            page_query += " LIMIT %s OFFSET %s"
            page_params.extend([params["length"], params["start"]])
        elif params["start"]:
            page_query += " LIMIT 18446744073709551615 OFFSET %s"
            page_params.append(params["start"])

        cur.execute(page_query, page_params)
        page_rows = _dictfetchall(cur)

    if not page_rows:
        return total, []

    latest_ids = [row["latest_id"] for row in page_rows if row.get("latest_id")]
    placeholders = ",".join(["%s"] * len(latest_ids))
    order_map = {row["latest_id"]: idx for idx, row in enumerate(page_rows)}

    detail_sql = f"""
        SELECT
            ids.id AS latest_id,
            ids.po_form_unique_id AS form_main_unique_id,
            ids.consignee_unique_id,
            ids.unique_id AS ins_unique_id,
            COALESCE(ids.po_num, pf.po_num, '') AS po_num,
            COALESCE(ids.po_date, pf.po_date, '') AS po_date,
            COALESCE(team.staff_name, team.user_name, NULLIF(TRIM(ids.team_mem), ''), NULLIF(TRIM(icm.team_mem), ''), NULLIF(TRIM(idm.team_mem), ''), NULLIF(TRIM(c.team_mem), ''), '') AS followed_by,
            COALESCE(dep_sub.ledger_name, dep.ledger_name, dep.department, dep_pf.department, icm.ledger_name, pf.department, '') AS ledger_name,
            COALESCE(cd.district_name, c.con_district, '') AS ledger_city,
            COALESCE(cs.state_name, c.con_state_name, '') AS ledger_state,
            ids.invoice_no,
            DATE_FORMAT(ids.invoice_date, '%%d-%%m-%%Y') AS invoice_date,
            ids.dc_number,
            COALESCE(icm.dc_date, '') AS dc_date,
            ids.documents_type AS dc_status,
            ids.dc_file AS dc_file,
            CASE
                WHEN COALESCE(NULLIF(TRIM(ids.documents_type1), ''), '0') <> '0' THEN ids.documents_type1
                WHEN {dc_required_sql} = '1' AND COALESCE(NULLIF(TRIM(ids.documents_type), ''), '0') <> '0' THEN ids.documents_type
                ELSE ids.documents_type1
            END AS ir_status,
            CASE
                WHEN COALESCE(ids.ir_file, '') <> '' THEN ids.ir_file
                WHEN {dc_required_sql} = '1' THEN COALESCE(ids.dc_file, '')
                ELSE ids.ir_file
            END AS ir_file,
            COALESCE((
                SELECT DATE_FORMAT(
                    CASE
                        WHEN {dc_required_sql} = '1' THEN dd.dc_pod_date
                        ELSE dd.ir_pod_date
                    END,
                    '%%d-%%m-%%Y'
                )
                FROM dc_ir_doc_dispatch_details dd
                WHERE dd.invoice_no = ids.invoice_no
                  AND dd.dc_number = ids.dc_number
                  AND dd.consignee_unique_id = ids.consignee_unique_id
                  AND dd.is_delete = 0
                LIMIT 1
            ), '') AS ir_pod_date,
            {dc_required_sql} AS dc_required,
            COALESCE(ids.document_verification_status, '0') AS document_verification_status
        FROM installation_details_sublist ids
        LEFT JOIN installation_details idm
          ON idm.po_form_unique_id = ids.po_form_unique_id
         AND idm.consignee_unique_id = ids.consignee_unique_id
         AND idm.invoice_no = ids.invoice_no
         AND idm.dc_number = ids.dc_number
         AND idm.is_delete = 0
        LEFT JOIN invoice_creation_main icm
          ON icm.form_main_unique_id = ids.po_form_unique_id
         AND icm.consignee_unique_id = ids.consignee_unique_id
         AND icm.invoice_no = ids.invoice_no
         AND icm.dc_number = ids.dc_number
         AND icm.is_delete = 0
        LEFT JOIN po_form pf
          ON ids.po_form_unique_id = pf.unique_id
         AND pf.is_delete = 0
        LEFT JOIN consignee_details_sub c
          ON c.unique_id = ids.consignee_unique_id
         AND c.is_delete = 0
        LEFT JOIN district_creation cd
          ON cd.unique_id = c.con_district
         AND cd.is_delete = 0
        LEFT JOIN state_creation cs
          ON cs.unique_id = c.con_state_name
         AND cs.is_delete = 0
        LEFT JOIN department_creation dep
          ON dep.unique_id = icm.ledger_name
         AND dep.is_delete = 0
        LEFT JOIN department_creation_sublist dep_sub
          ON dep_sub.unique_id = icm.ledger_name
         AND dep_sub.is_delete = 0
        LEFT JOIN department_creation dep_pf
          ON dep_pf.unique_id = pf.department
         AND dep_pf.is_delete = 0
        LEFT JOIN user team
          ON (
                team.staff_id = COALESCE(NULLIF(TRIM(ids.team_mem), ''), NULLIF(TRIM(icm.team_mem), ''), NULLIF(TRIM(idm.team_mem), ''), NULLIF(TRIM(c.team_mem), ''))
                OR team.unique_id = COALESCE(NULLIF(TRIM(ids.team_mem), ''), NULLIF(TRIM(icm.team_mem), ''), NULLIF(TRIM(idm.team_mem), ''), NULLIF(TRIM(c.team_mem), ''))
                OR team.user_name = COALESCE(NULLIF(TRIM(ids.team_mem), ''), NULLIF(TRIM(icm.team_mem), ''), NULLIF(TRIM(idm.team_mem), ''), NULLIF(TRIM(c.team_mem), ''))
             )
         AND team.is_delete = 0
        WHERE ids.id IN ({placeholders})
    """

    with connection.cursor() as cur:
        cur.execute(detail_sql, latest_ids)
        rows = _dictfetchall(cur)

    rows.sort(key=lambda row: order_map.get(row.get("latest_id"), 10**9))
    rows = _unique_rows(
        rows,
        lambda row: (
            row.get("latest_id"),
            row.get("consignee_unique_id", ""),
            row.get("invoice_no", ""),
            row.get("dc_number", ""),
        ),
    )

    data = []
    for idx, row in enumerate(rows, start=params["start"] + 1):
        dc_status_raw = str(row.get("dc_status") or "").strip()
        ir_status_raw = str(row.get("ir_status") or "").strip()
        dc_status = "Yes" if dc_status_raw and dc_status_raw != "0" else "-"
        ir_status = "NA" if ir_status_raw in {"", "0", "undefined"} else "YES"
        data.append(
            {
                "id": f"{row.get('consignee_unique_id', '')}|{row.get('invoice_no', '')}|{row.get('dc_number', '')}|{row.get('ins_unique_id', '')}",
                "verificationUniqueId": "",
                "consigneeUniqueId": row.get("consignee_unique_id", ""),
                "formMainUniqueId": row.get("form_main_unique_id", ""),
                "insUniqueId": row.get("ins_unique_id", ""),
                "poNo": row.get("po_num", ""),
                "poDate": _fmt_date(row.get("po_date")),
                "followedBy": row.get("followed_by", ""),
                "ledgerName": row.get("ledger_name", ""),
                "ledgerCity": row.get("ledger_city", ""),
                "ledgerState": row.get("ledger_state", ""),
                "invoiceNo": row.get("invoice_no", ""),
                "invoiceDate": _fmt_date(row.get("invoice_date")),
                "dcNo": row.get("dc_number", ""),
                "dcDate": _fmt_date(row.get("dc_date")),
                "dcRecStatus": dc_status,
                "dcSignAtt": _bool_file(row.get("dc_file")),
                "dcFileUrl": _installation_file_url(row.get("dc_file")),
                "irRecStatus": ir_status,
                "irSignAtt": _bool_file(row.get("ir_file")),
                "irFileUrl": _installation_file_url(row.get("ir_file")),
                "irPodDate": row.get("ir_pod_date", "") or "",
                "snrRecStatus": row.get("snr_status", ""),
                "snrSignAtt": _bool_file(row.get("snr_file")),
                "snrFileUrl": _installation_file_url(row.get("snr_file")),
                "snrPodDate": row.get("snr_pod_date", "") or "",
                "status": _status_label(row.get("document_verification_status")),
                "canDelete": bool(row.get("consignee_unique_id")) and bool(row.get("invoice_no")) and bool(row.get("dc_number")) and bool(row.get("ins_unique_id")),
                "rejectReason": "",
                "s_no": idx,
            }
        )
    return total, data


def _get_pending_snr_rows(params: dict):
    conditions = [
        "ids.is_delete = 0",
        "COALESCE(ids.document_verification_status, '0') = '0'",
        "COALESCE(ids.documents_type1, '0') = '0'",
        "COALESCE(dd.dc_dispatch_mode, '') <> ''",
        "COALESCE(dd.snr_dispatch_mode, '') <> ''",
        "dd.is_delete = 0",
    ]
    sql_params: list[object] = []
    if params.get("company_id"):
        conditions.append("p.sess_company_id = %s")
        sql_params.append(params["company_id"])

    if params["from_date"] and params["to_date"]:
        date_column = {
            "po_date": "DATE(COALESCE(ids.po_date, p.po_date))",
            "invoice_date": "DATE(COALESCE(ids.invoice_date, icm.invoice_date))",
            "dc_date": "DATE(COALESCE(dl.dc_date, icm.dc_date))",
        }.get(_date_column(params["opt"]).split(".")[-1], "DATE(COALESCE(ids.po_date, p.po_date))")
        conditions.append(f"{date_column} BETWEEN %s AND %s")
        sql_params.extend([params["from_date"], params["to_date"]])

    if params["followed_by"] and params["followed_by"].lower() != "all":
        conditions.append(
            "(COALESCE(u.staff_name, u.user_name, ids.team_mem, icm.team_mem, idm.team_mem, c.team_mem, '') = %s OR COALESCE(ids.team_mem, icm.team_mem, idm.team_mem, c.team_mem, '') = %s)"
        )
        sql_params.extend([params["followed_by"], params["followed_by"]])

    if params["search"]:
        like = f"%{params['search']}%"
        conditions.append(
            """(
                COALESCE(ids.po_num, '') LIKE %s
                OR COALESCE(d.ledger_name, '') LIKE %s
                OR COALESCE(ids.invoice_no, '') LIKE %s
                OR COALESCE(ids.dc_number, '') LIKE %s
                OR COALESCE(u.staff_name, u.user_name, ids.team_mem, icm.team_mem, idm.team_mem, c.team_mem, '') LIKE %s
                OR COALESCE(cd.district_name, c.con_district, '') LIKE %s
                OR COALESCE(cs.state_name, c.con_state_name, '') LIKE %s
            )"""
        )
        sql_params.extend([like, like, like, like, like, like, like])

    where_sql = " AND ".join(conditions)
    from_sql = f"""
        FROM installation_details_sublist ids
        LEFT JOIN installation_details idm
          ON idm.po_form_unique_id = ids.po_form_unique_id
         AND idm.consignee_unique_id = ids.consignee_unique_id
         AND idm.invoice_no = ids.invoice_no
         AND idm.dc_number = ids.dc_number
         AND idm.is_delete = 0
        LEFT JOIN invoice_creation_main icm
          ON icm.form_main_unique_id = ids.po_form_unique_id
         AND icm.consignee_unique_id = ids.consignee_unique_id
         AND icm.invoice_no = ids.invoice_no
         AND icm.dc_number = ids.dc_number
         AND icm.is_delete = 0
        LEFT JOIN po_form p
          ON p.unique_id = ids.po_form_unique_id
         AND p.is_delete = 0
        LEFT JOIN consignee_details_sub c
          ON c.unique_id = ids.consignee_unique_id
         AND c.is_delete = 0
        LEFT JOIN district_creation cd
          ON cd.unique_id = c.con_district
         AND cd.is_delete = 0
        LEFT JOIN state_creation cs
          ON cs.unique_id = c.con_state_name
         AND cs.is_delete = 0
        LEFT JOIN department_creation d
          ON d.unique_id = p.department
         AND d.is_delete = 0
        LEFT JOIN dispatch_list dl
          ON dl.invoice_no = ids.invoice_no
         AND dl.dc_number = ids.dc_number
         AND dl.consignee_unique_id = ids.consignee_unique_id
         AND dl.is_delete = 0
        LEFT JOIN dc_ir_doc_dispatch_details dd
          ON dd.po_form_unique_id = ids.po_form_unique_id
         AND dd.consignee_unique_id = ids.consignee_unique_id
         AND dd.invoice_no = ids.invoice_no
         AND dd.dc_number = ids.dc_number
         AND dd.is_delete = 0
        LEFT JOIN `user` u
          ON (
                u.staff_id = COALESCE(ids.team_mem, icm.team_mem, idm.team_mem, c.team_mem)
                OR u.unique_id = COALESCE(ids.team_mem, icm.team_mem, idm.team_mem, c.team_mem)
                OR u.user_name = COALESCE(ids.team_mem, icm.team_mem, idm.team_mem, c.team_mem)
             )
         AND u.is_delete = 0
        WHERE {where_sql}
    """

    count_sql = f"""
        SELECT COUNT(*)
        FROM (
            SELECT 1
            {from_sql}
            GROUP BY ids.po_form_unique_id, ids.consignee_unique_id, ids.invoice_no, ids.dc_number
        ) q
    """

    page_sql = f"""
        SELECT
            MAX(ids.id) AS latest_id,
            ids.po_form_unique_id,
            ids.consignee_unique_id,
            ids.invoice_no,
            ids.dc_number
        {from_sql}
        GROUP BY ids.po_form_unique_id, ids.consignee_unique_id, ids.invoice_no, ids.dc_number
        ORDER BY MAX(ids.id) DESC
    """

    with connection.cursor() as cur:
        cur.execute(count_sql, sql_params)
        total = cur.fetchone()[0] or 0

        page_query = page_sql
        page_params = list(sql_params)
        if params["length"] != -1:
            page_query += " LIMIT %s OFFSET %s"
            page_params.extend([params["length"], params["start"]])
        elif params["start"]:
            page_query += " LIMIT 18446744073709551615 OFFSET %s"
            page_params.append(params["start"])

        cur.execute(page_query, page_params)
        page_rows = _dictfetchall(cur)

    if not page_rows:
        return total, []

    latest_ids = [row["latest_id"] for row in page_rows if row.get("latest_id")]
    placeholders = ",".join(["%s"] * len(latest_ids))
    order_map = {row["latest_id"]: idx for idx, row in enumerate(page_rows)}

    detail_sql = f"""
        SELECT
            ids.id AS latest_id,
            ids.po_form_unique_id AS form_main_unique_id,
            ids.consignee_unique_id,
            ids.unique_id AS ins_unique_id,
            COALESCE(ids.po_num, p.po_num, '') AS po_num,
            COALESCE(ids.po_date, p.po_date, '') AS po_date,
            COALESCE(u.staff_name, u.user_name, ids.team_mem, icm.team_mem, idm.team_mem, c.team_mem, '') AS followed_by,
            COALESCE(d.department, '') AS ledger_name,
            COALESCE(cd.district_name, c.con_district, '') AS ledger_city,
            COALESCE(cs.state_name, c.con_state_name, '') AS ledger_state,
            ids.invoice_no,
            COALESCE(ids.invoice_date, icm.invoice_date, '') AS invoice_date,
            ids.dc_number,
            COALESCE(dl.dc_date, icm.dc_date) AS dc_date,
            COALESCE(ids.documents_type, '0') AS dc_status,
            COALESCE(ids.documents_type1, '0') AS ir_status,
            COALESCE(ids.documents_type2, '0') AS snr_status,
            COALESCE(ids.dc_file, idm.dc_file, '') AS dc_file,
            COALESCE(ids.ir_file, idm.ir_file, '') AS ir_file,
            COALESCE(ids.snr_file, idm.snr_file, '') AS snr_file,
            COALESCE((
                SELECT DATE_FORMAT(dispatch.snr_pod_date, '%%d-%%m-%%Y')
                FROM dc_ir_doc_dispatch_details dispatch
                WHERE dispatch.invoice_no = ids.invoice_no
                  AND dispatch.dc_number = ids.dc_number
                  AND dispatch.consignee_unique_id = ids.consignee_unique_id
                  AND dispatch.is_delete = 0
                LIMIT 1
            ), '') AS snr_pod_date,
            COALESCE(ids.document_verification_status, '0') AS document_verification_status
        FROM installation_details_sublist ids
        LEFT JOIN installation_details idm
          ON idm.po_form_unique_id = ids.po_form_unique_id
         AND idm.consignee_unique_id = ids.consignee_unique_id
         AND idm.invoice_no = ids.invoice_no
         AND idm.dc_number = ids.dc_number
         AND idm.is_delete = 0
        LEFT JOIN invoice_creation_main icm
          ON icm.form_main_unique_id = ids.po_form_unique_id
         AND icm.consignee_unique_id = ids.consignee_unique_id
         AND icm.invoice_no = ids.invoice_no
         AND icm.dc_number = ids.dc_number
         AND icm.is_delete = 0
        LEFT JOIN po_form p
          ON p.unique_id = ids.po_form_unique_id
         AND p.is_delete = 0
        LEFT JOIN consignee_details_sub c
          ON c.unique_id = ids.consignee_unique_id
         AND c.is_delete = 0
        LEFT JOIN district_creation cd
          ON cd.unique_id = c.con_district
         AND cd.is_delete = 0
        LEFT JOIN state_creation cs
          ON cs.unique_id = c.con_state_name
         AND cs.is_delete = 0
        LEFT JOIN department_creation d
          ON d.unique_id = p.department
         AND d.is_delete = 0
        LEFT JOIN dispatch_list dl
          ON dl.invoice_no = ids.invoice_no
         AND dl.dc_number = ids.dc_number
         AND dl.consignee_unique_id = ids.consignee_unique_id
         AND dl.is_delete = 0
        LEFT JOIN `user` u
          ON (
                u.staff_id = COALESCE(ids.team_mem, icm.team_mem, idm.team_mem, c.team_mem)
                OR u.unique_id = COALESCE(ids.team_mem, icm.team_mem, idm.team_mem, c.team_mem)
                OR u.user_name = COALESCE(ids.team_mem, icm.team_mem, idm.team_mem, c.team_mem)
             )
         AND u.is_delete = 0
        WHERE ids.id IN ({placeholders})
    """

    with connection.cursor() as cur:
        cur.execute(detail_sql, latest_ids)
        rows = _dictfetchall(cur)

    rows.sort(key=lambda row: order_map.get(row.get("latest_id"), 10**9))
    rows = _unique_rows(
        rows,
        lambda row: (
            row.get("latest_id"),
            row.get("consignee_unique_id", ""),
            row.get("invoice_no", ""),
            row.get("dc_number", ""),
        ),
    )

    data = []
    for idx, row in enumerate(rows, start=params["start"] + 1):
        dc_status_raw = str(row.get("dc_status") or "").strip()
        snr_status_raw = str(row.get("snr_status") or "").strip()
        data.append(
            {
                "id": f"{row.get('consignee_unique_id', '')}|{row.get('invoice_no', '')}|{row.get('dc_number', '')}|{row.get('ins_unique_id', '')}",
                "verificationUniqueId": "",
                "consigneeUniqueId": row.get("consignee_unique_id", ""),
                "formMainUniqueId": row.get("form_main_unique_id", ""),
                "insUniqueId": row.get("ins_unique_id", ""),
                "poNo": row.get("po_num", ""),
                "poDate": _fmt_date(row.get("po_date")),
                "followedBy": row.get("followed_by", ""),
                "ledgerName": row.get("ledger_name", ""),
                "ledgerCity": row.get("ledger_city", ""),
                "ledgerState": row.get("ledger_state", ""),
                "invoiceNo": row.get("invoice_no", ""),
                "invoiceDate": _fmt_date(row.get("invoice_date")),
                "dcNo": row.get("dc_number", ""),
                "dcDate": _fmt_date(row.get("dc_date")),
                "dcRecStatus": "Yes" if dc_status_raw and dc_status_raw != "0" else "-",
                "dcSignAtt": _bool_file(row.get("dc_file")),
                "dcFileUrl": _installation_file_url(row.get("dc_file")),
                "irRecStatus": row.get("ir_status", ""),
                "irSignAtt": _bool_file(row.get("ir_file")),
                "irFileUrl": _installation_file_url(row.get("ir_file")),
                "irPodDate": row.get("ir_pod_date", "") or "",
                "snrRecStatus": "Yes" if snr_status_raw and snr_status_raw != "0" else "-",
                "snrSignAtt": _bool_file(row.get("snr_file")),
                "snrFileUrl": _installation_file_url(row.get("snr_file")),
                "snrPodDate": row.get("snr_pod_date", "") or "",
                "status": _status_label(row.get("document_verification_status")),
                "canDelete": bool(row.get("consignee_unique_id")) and bool(row.get("invoice_no")) and bool(row.get("dc_number")) and bool(row.get("ins_unique_id")),
                "rejectReason": "",
                "s_no": idx,
            }
        )
    return total, data


def _get_pending_rows(params: dict):
    if params["pending_type"] == "ir":
        return _get_pending_ir_rows(params)
    if params["pending_type"] == "snr":
        return _get_pending_snr_rows(params)

    base_sql, select_sql, sql_params = _build_pending_query(params)
    total, rows = _paginate_rows(base_sql, select_sql, sql_params, params["start"], params["length"])
    rows = _unique_rows(
        rows,
        lambda row: (
            row.get("consignee_unique_id", ""),
            row.get("invoice_no", ""),
            row.get("dc_number", ""),
        ),
    )
    data = []
    for idx, row in enumerate(rows, start=params["start"] + 1):
        if params["pending_type"] == "ir":
            dc_status_raw = str(row.get("dc_status") or "").strip()
            ir_status_raw = str(row.get("ir_status") or "").strip()
            dc_status = "Yes" if dc_status_raw and dc_status_raw != "0" else "-"
            ir_status = "NA" if ir_status_raw in {"", "0", "undefined"} else "YES"
        else:
            dc_status = row.get("dc_status", "")
            ir_status = row.get("ir_status", "")
            snr_status_raw = str(row.get("snr_status") or "").strip()
            dc_status = "Yes" if str(dc_status or "").strip() and str(dc_status).strip() != "0" else "-"
            ir_status = row.get("ir_status", "")
            snr_status = "Yes" if snr_status_raw and snr_status_raw != "0" else "-"
        data.append(
            {
                "id": f"{row.get('consignee_unique_id', '')}|{row.get('invoice_no', '')}|{row.get('dc_number', '')}|{row.get('ins_unique_id', '')}",
                "verificationUniqueId": "",
                "consigneeUniqueId": row.get("consignee_unique_id", ""),
                "formMainUniqueId": row.get("form_main_unique_id", ""),
                "insUniqueId": row.get("ins_unique_id", ""),
                "poNo": row.get("po_num", ""),
                "poDate": _fmt_date(row.get("po_date")),
                "followedBy": row.get("followed_by", ""),
                "ledgerName": row.get("ledger_name", ""),
                "ledgerCity": row.get("ledger_city", ""),
                "ledgerState": row.get("ledger_state", ""),
                "invoiceNo": row.get("invoice_no", ""),
                "invoiceDate": _fmt_date(row.get("invoice_date")),
                "dcNo": row.get("dc_number", ""),
                "dcDate": _fmt_date(row.get("dc_date")),
                "dcRecStatus": dc_status,
                "dcSignAtt": _bool_file(row.get("dc_file")),
                "dcFileUrl": _installation_file_url(row.get("dc_file")),
                "irRecStatus": ir_status,
                "irSignAtt": _bool_file(row.get("ir_file")),
                "irFileUrl": _installation_file_url(row.get("ir_file")),
                "irPodDate": row.get("ir_pod_date", "") or "",
                "snrRecStatus": snr_status if params["pending_type"] == "snr" else row.get("snr_status", ""),
                "snrSignAtt": _bool_file(row.get("snr_file")),
                "snrFileUrl": _installation_file_url(row.get("snr_file")),
                "snrPodDate": row.get("snr_pod_date", "") or "",
                "status": _status_label(row.get("document_verification_status")),
                "canDelete": bool(row.get("consignee_unique_id")) and bool(row.get("invoice_no")) and bool(row.get("dc_number")) and bool(row.get("ins_unique_id")),
                "rejectReason": "",
                "s_no": idx,
            }
        )
    return total, data


def _get_saved_rows(params: dict, statuses: list[str]):
    if statuses == ["2"]:
        base_sql, select_sql, sql_params = _build_verified_query(params)
    else:
        base_sql, select_sql, sql_params = _build_saved_query(params, statuses)
    total, rows = _paginate_rows(base_sql, select_sql, sql_params, params["start"], params["length"])
    rows = _unique_rows(
        rows,
        lambda row: (
            row.get("verification_unique_id", ""),
            row.get("consignee_unique_id", ""),
            row.get("invoice_no", ""),
            row.get("dc_number", ""),
        ),
    )
    data = []
    for idx, row in enumerate(rows, start=params["start"] + 1):
        # Allow delete when:
        # 1. verification_unique_id exists and no bill, OR
        # 2. DC/invoice/consignee keys exist and no bill (can match by dc_number)
        has_bill = bool(str(row.get("bill_no", "")).strip())
        has_verification_id = bool(row.get("verification_unique_id"))
        has_dc_keys = bool(row.get("dc_number")) and bool(row.get("invoice_no")) and bool(row.get("consignee_unique_id"))
        can_delete = not has_bill and (has_verification_id or has_dc_keys)
        
        data.append(
            {
                "id": row.get("verification_unique_id", ""),
                "verificationUniqueId": row.get("verification_unique_id", ""),
                "consigneeUniqueId": row.get("consignee_unique_id", ""),
                "formMainUniqueId": row.get("form_main_unique_id", ""),
                "insUniqueId": row.get("ins_unique_id", ""),
                "poNo": row.get("po_num", ""),
                "poDate": _fmt_date(row.get("po_date")),
                "followedBy": row.get("followed_by", ""),
                "ledgerName": row.get("ledger_name", ""),
                "ledgerCity": row.get("ledger_city", ""),
                "ledgerState": row.get("ledger_state", ""),
                "invoiceNo": row.get("invoice_no", ""),
                "invoiceDate": _fmt_date(row.get("invoice_date")),
                "dcNo": row.get("dc_number", ""),
                "dcDate": _fmt_date(row.get("dc_date")),
                "dcSignedDoc": _bool_file(row.get("dc_file")),
                "dcFileUrl": _installation_file_url(row.get("dc_file")),
                "installSignedReport": _bool_file(row.get("snr_file") or row.get("ir_file") or row.get("dc_file")),
                "installSignedReportUrl": _installation_file_url(row.get("snr_file") or row.get("ir_file") or row.get("dc_file")),
                "status": _status_label(row.get("status_app")),
                "rejectReason": row.get("reject_reason", ""),
                "canDelete": can_delete,
                "s_no": idx,
            }
        )
    return total, data


def _get_detail_base(consignee_unique_id: str, invoice_no: str, dc_number: str, ins_unique_id: str):
    with connection.cursor() as cur:
        cur.execute(
            f"""
            SELECT
                ids.po_form_unique_id AS form_main_unique_id,
                ids.consignee_unique_id,
                ids.unique_id AS ins_unique_id,
                COALESCE(ids.po_num, p.po_num, '') AS po_num,
                COALESCE(ids.po_date, p.po_date, '') AS po_date,
                ids.invoice_no,
                COALESCE(ids.invoice_date, icm.invoice_date, '') AS invoice_date,
                ids.dc_number,
                COALESCE(d.dc_date, icm.dc_date, '') AS dc_date,
                COALESCE(ids.dc_received_sts, i.dc_received_sts, '') AS dc_status,
                COALESCE(ids.ir_rec_status, i.ir_rec_status, '') AS ir_status,
                COALESCE(ids.snr_rec_status, i.snr_rec_status, '') AS snr_status,
                COALESCE(p.file_name, '') AS po_file,
                COALESCE(ids.dc_file, '') AS dc_file,
                COALESCE(ids.ir_file, '') AS ir_file,
                COALESCE(ids.snr_file, '') AS snr_file,
                COALESCE(dep.department, dep_sub.ledger_name, icm.ledger_name, '') AS ledger_name,
                COALESCE(p.bill_address, '') AS bill_address,
                COALESCE(p.contact_number, '') AS customer_phone,
                COALESCE(p.email, '') AS customer_email,
                COALESCE(cust_district.district_name, p.district, '') AS customer_district,
                COALESCE(cust_state.state_name, p.state_name, '') AS customer_state,
                COALESCE(c.con_contact_name, '') AS consignee_name,
                COALESCE(c.con_address, '') AS consignee_address,
                COALESCE(c.con_contact_number, '') AS consignee_phone,
                COALESCE(cons_district.district_name, c.con_district, '') AS consignee_district,
                COALESCE(cons_state.state_name, c.con_state_name, '') AS consignee_state,
                COALESCE(c.con_pincode, '') AS consignee_pincode,
                COALESCE(d.delivery_date, '') AS delivery_date,
                COALESCE(dd.dc_pod_no, '') AS dc_pod_no,
                COALESCE(dd.dc_pod_date, '') AS dc_pod_date,
                COALESCE(dd.ir_pod_no, '') AS ir_pod_no,
                COALESCE(dd.ir_pod_date, '') AS ir_pod_date,
                COALESCE(dd.snr_pod_no, '') AS snr_pod_no,
                COALESCE(dd.snr_pod_date, '') AS snr_pod_date,
                COALESCE(i.documents_type, '') AS documents_type,
                COALESCE(i.documents_type1, '') AS documents_type1,
                COALESCE(i.documents_type2, '') AS documents_type2,
                {_dc_required_sql("i", "ids.po_form_unique_id")} AS dc_required,
                COALESCE(i.snr_verify_status, 0) AS snr_verify_status
            FROM installation_details_sublist ids
            LEFT JOIN invoice_creation_main icm
              ON icm.form_main_unique_id = ids.po_form_unique_id
             AND icm.consignee_unique_id = ids.consignee_unique_id
             AND icm.invoice_no = ids.invoice_no
             AND icm.dc_number = ids.dc_number
             AND icm.is_delete = 0
            LEFT JOIN po_form p
              ON p.unique_id = ids.po_form_unique_id
             AND p.is_delete = 0
            LEFT JOIN district_creation cust_district
              ON cust_district.unique_id = p.district
             AND cust_district.is_delete = 0
            LEFT JOIN state_creation cust_state
              ON cust_state.unique_id = p.state_name
             AND cust_state.is_delete = 0
            LEFT JOIN consignee_details_sub c
              ON c.unique_id = ids.consignee_unique_id
             AND c.is_delete = 0
            LEFT JOIN district_creation cons_district
              ON cons_district.unique_id = c.con_district
             AND cons_district.is_delete = 0
            LEFT JOIN state_creation cons_state
              ON cons_state.unique_id = c.con_state_name
             AND cons_state.is_delete = 0
            LEFT JOIN dispatch_list d
              ON d.invoice_no = ids.invoice_no
             AND d.dc_number = ids.dc_number
             AND d.consignee_unique_id = ids.consignee_unique_id
             AND d.is_delete = 0
            LEFT JOIN dc_ir_doc_dispatch_details dd
              ON dd.po_form_unique_id = ids.po_form_unique_id
             AND dd.consignee_unique_id = ids.consignee_unique_id
             AND dd.invoice_no = ids.invoice_no
             AND dd.dc_number = ids.dc_number
             AND dd.is_delete = 0
            LEFT JOIN installation_details i
              ON i.po_form_unique_id = ids.po_form_unique_id
             AND i.consignee_unique_id = ids.consignee_unique_id
             AND i.invoice_no = ids.invoice_no
             AND i.dc_number = ids.dc_number
             AND i.is_delete = 0
            LEFT JOIN department_creation dep
              ON dep.unique_id = icm.ledger_name
             AND dep.is_delete = 0
            LEFT JOIN department_creation_sublist dep_sub
              ON dep_sub.unique_id = icm.ledger_name
             AND dep_sub.is_delete = 0
            WHERE ids.consignee_unique_id = %s
              AND ids.invoice_no = %s
              AND ids.dc_number = %s
              AND ids.unique_id = %s
              AND ids.is_delete = 0
            LIMIT 1
            """,
            [consignee_unique_id, invoice_no, dc_number, ins_unique_id],
        )
        rows = _dictfetchall(cur)
    return rows[0] if rows else None


def _get_existing_verification(consignee_unique_id: str, invoice_no: str, dc_number: str, ins_unique_id: str):
    with connection.cursor() as cur:
        cur.execute(
            """
            SELECT *
            FROM sign_doc_verification_detail
            WHERE con_unique_id = %s
              AND invoice_no = %s
              AND dc_number = %s
              AND ins_unique_id = %s
              AND is_delete = 0
            ORDER BY id DESC
            LIMIT 1
            """,
            [consignee_unique_id, invoice_no, dc_number, ins_unique_id],
        )
        rows = _dictfetchall(cur)
    return rows[0] if rows else None


def _get_item_rows(form_main_unique_id: str):
    with connection.cursor() as cur:
        cur.execute(
            """
            SELECT item_code, product, qty, total_value
            FROM product_details_sub
            WHERE form_main_unique_id = %s
              AND is_delete = 0
            ORDER BY id ASC
            """,
            [form_main_unique_id],
        )
        rows = _dictfetchall(cur)
    items = []
    for idx, row in enumerate(rows, start=1):
        try:
            value = float(row.get("total_value") or 0)
        except (TypeError, ValueError):
            value = 0.0
        try:
            qty = int(float(row.get("qty") or 0))
        except (TypeError, ValueError):
            qty = 0
        items.append(
            {
                "id": idx,
                "itemName": row.get("item_code", ""),
                "itemDesc": row.get("product", ""),
                "dcQty": qty,
                "invoiceValue": value,
            }
        )
    return items


class SignDocVerificationListView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        params = _parse_list_params(request)
        params["company_id"] = request_company_id(request)
        if params["tab"] == "mismatch":
            total, data = _get_saved_rows(params, ["1"])
        elif params["tab"] == "verified":
            total, data = _get_saved_rows(params, ["2"])
        else:
            total, data = _get_pending_rows(params)
        return Response(
            {
                "draw": params["draw"],
                "recordsTotal": total,
                "recordsFiltered": total,
                "data": data,
            }
        )


class SignDocVerificationDetailView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request, consignee_unique_id):
        invoice_no = str(request.query_params.get("invoice_no", "")).strip()
        dc_number = str(request.query_params.get("dc_number", "")).strip()
        ins_unique_id = str(request.query_params.get("ins_unique_id", "")).strip()
        if not (consignee_unique_id and invoice_no and dc_number and ins_unique_id):
            return Response({"status": False, "message": "Missing record identifiers."}, status=status.HTTP_400_BAD_REQUEST)

        base = _get_detail_base(consignee_unique_id, invoice_no, dc_number, ins_unique_id)
        if base is None:
            return Response({"status": False, "message": "Record not found."}, status=status.HTTP_404_NOT_FOUND)

        existing = _get_existing_verification(consignee_unique_id, invoice_no, dc_number, ins_unique_id)
        status_label = _status_label((existing or {}).get("status_app", "0"))
        customer_address = ", ".join(
            [part for part in [base.get("bill_address", ""), base.get("customer_district", ""), base.get("customer_state", "")] if part]
        )
        consignee_city = ", ".join(
            [part for part in [base.get("consignee_district", ""), base.get("consignee_state", "")] if part]
        )
        if base.get("consignee_pincode"):
            consignee_city = f"{consignee_city} - {base.get('consignee_pincode')}" if consignee_city else str(base.get("consignee_pincode"))

        data = {
            "verificationUniqueId": (existing or {}).get("unique_id", ""),
            "consigneeUniqueId": consignee_unique_id,
            "formMainUniqueId": base.get("form_main_unique_id", ""),
            "insUniqueId": base.get("ins_unique_id", ""),
            "customerName": base.get("ledger_name", ""),
            "customerAddress": customer_address,
            "customerPhone": base.get("customer_phone") or "--",
            "customerEmail": base.get("customer_email") or "--",
            "consigneeName": base.get("consignee_name", ""),
            "consigneeAddress": base.get("consignee_address", ""),
            "consigneePhone": base.get("consignee_phone") or "--",
            "consigneeCity": consignee_city,
            "poNumber": base.get("po_num", ""),
            "poDate": _fmt_date(base.get("po_date")),
            "invoiceNo": base.get("invoice_no", ""),
            "invoiceDate": _fmt_date(base.get("invoice_date")),
            "dcNumber": base.get("dc_number", ""),
            "dcDate": _fmt_date(base.get("dc_date")),
            "deliveryDate": _fmt_date(base.get("delivery_date"), "%d/%m/%Y"),
            "dcPodNo": base.get("dc_pod_no", ""),
            "dcPodDate": _fmt_date(base.get("dc_pod_date")),
            "irPodNo": base.get("ir_pod_no", ""),
            "irPodDate": _fmt_date(base.get("ir_pod_date")),
            "snrPodNo": base.get("snr_pod_no", ""),
            "snrPodDate": _fmt_date(base.get("snr_pod_date")),
            "poAttachment": base.get("po_file") or None,
            "dcSignedDocument": base.get("dc_file") or None,
            "installationSignedReport": (base.get("snr_file") or base.get("ir_file")) or None,
            "items": _get_item_rows(base.get("form_main_unique_id", "")),
            "dcReceivedStatus": (existing or {}).get("dc_received_status") or base.get("dc_status", "") or "Pending",
            "dcSignedDate": _fmt_date((existing or {}).get("dc_signed_date"), "%Y-%m-%d"),
            "irReceivedStatus": (existing or {}).get("ir_status") or base.get("ir_status", "") or ((existing or {}).get("dc_received_status") if str(base.get("dc_required", "") or "").strip() == "1" else "") or "Pending",
            "irSignedDate": _fmt_date((existing or {}).get("ir_signed_date") or ((existing or {}).get("dc_signed_date") if str(base.get("dc_required", "") or "").strip() == "1" else ""), "%Y-%m-%d"),
            "snrReceivedStatus": (existing or {}).get("snr_status") or base.get("snr_status", "") or "Pending",
            "snrSignedDate": _fmt_date((existing or {}).get("snr_signed_date"), "%Y-%m-%d"),
            "hoReceivedDate": _fmt_date((existing or {}).get("inv_verify_approved_date"), "%Y-%m-%d"),
            "docVerification": "Mismatch/Rejected" if status_label == "Mismatch" else status_label,
            "primaryProductForBg": (existing or {}).get("po_product_name", ""),
            "processedWithBg": str((existing or {}).get("with_bg", "")).lower() == "on",
            "processedWithoutBg": str((existing or {}).get("without_bg", "")).lower() == "on",
            "rejectReason": (existing or {}).get("reject_reason", ""),
            "docChn": (existing or {}).get("doc_chn", ""),
            "sts": base.get("documents_type", ""),
            "sts1": base.get("documents_type1", ""),
            "sts2": base.get("documents_type2", ""),
            "dcRequired": base.get("dc_required", ""),
            "snrVerifyStatus": int(base.get("snr_verify_status") or 0),
        }
        return Response({"status": True, "data": data})


class SignDocVerificationSaveView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        data = request.data
        form_main_unique_id = str(data.get("formMainUniqueId", "")).strip()
        consignee_unique_id = str(data.get("consigneeUniqueId", "")).strip()
        ins_unique_id = str(data.get("insUniqueId", "")).strip()
        po_num = str(data.get("poNumber", "")).strip()
        po_date = str(data.get("poDate", "")).strip()
        invoice_no = str(data.get("invoiceNo", "")).strip()
        invoice_date = str(data.get("invoiceDate", "")).strip()
        dc_number = str(data.get("dcNumber", "")).strip()
        verification_unique_id = str(data.get("verificationUniqueId", "")).strip()
        con_contact_name = str(data.get("consigneeName", "")).strip()
        con_address = str(data.get("consigneeAddress", "")).strip()
        dc_received_status = str(data.get("dcReceivedStatus", "")).strip()
        dc_signed_date = str(data.get("dcSignedDate", "")).strip()
        ir_status = str(data.get("irReceivedStatus", "")).strip()
        ir_signed_date = str(data.get("irSignedDate", "")).strip()
        snr_status = str(data.get("snrReceivedStatus", "")).strip()
        snr_signed_date = str(data.get("snrSignedDate", "")).strip()
        ho_received_date = str(data.get("hoReceivedDate", "")).strip()
        doc_verification = str(data.get("docVerification", "Pending")).strip()
        primary_product = str(data.get("primaryProductForBg", "")).strip()
        processed_with_bg = bool(data.get("processedWithBg"))
        processed_without_bg = bool(data.get("processedWithoutBg"))
        reject_reason = str(data.get("rejectReason", "")).strip()
        doc_chn = str(data.get("docChn", "")).strip()
        snr_verify_status = int(data.get("snrVerifyStatus") or 0)

        if not (form_main_unique_id and consignee_unique_id and ins_unique_id and invoice_no and dc_number):
            return Response({"status": False, "message": "Missing required record identifiers."}, status=status.HTTP_400_BAD_REQUEST)
        if doc_verification == "Mismatch/Rejected" and not reject_reason:
            return Response({"status": False, "message": "Reject reason is required."}, status=status.HTTP_400_BAD_REQUEST)

        status_code = _status_code_from_label(doc_verification)
        unique_id = verification_unique_id or _generate_unique_id()
        approved_by = str(
            getattr(request.user, "staff_name", "")
            or getattr(request.user, "user_name", "")
            or getattr(request.user, "staff_id", "")
            or request.user
        )
        approved_date = datetime.date.today().strftime("%Y-%m-%d")

        try:
            with transaction.atomic():
                with connection.cursor() as cur:
                    cur.execute(
                        """
                        SELECT id, unique_id
                        FROM sign_doc_verification_detail
                        WHERE con_unique_id = %s
                          AND invoice_no = %s
                          AND dc_number = %s
                          AND ins_unique_id = %s
                          AND is_delete = 0
                        ORDER BY id DESC
                        LIMIT 1
                        """,
                        [consignee_unique_id, invoice_no, dc_number, ins_unique_id],
                    )
                    existing = cur.fetchone()

                    if existing:
                        unique_id = verification_unique_id or existing[1]
                        cur.execute(
                            """
                            UPDATE sign_doc_verification_detail
                            SET form_main_unique_id = %s,
                                po_num = %s,
                                po_date = %s,
                                con_contact_name = %s,
                                con_address = %s,
                                con_unique_id = %s,
                                dc_received_status = %s,
                                dc_signed_date = %s,
                                ir_status = %s,
                                ir_signed_date = %s,
                                snr_status = %s,
                                snr_signed_date = %s,
                                po_product_name = %s,
                                with_bg = %s,
                                without_bg = %s,
                                invoice_no = %s,
                                dc_number = %s,
                                invoice_date = %s,
                                status_app = %s,
                                doc_chn = %s,
                                reject_reason = %s,
                                ins_unique_id = %s,
                                snr_verify_status = %s,
                                inv_verify_status = 1,
                                inv_verify_approvedby = %s,
                                inv_verify_approved_date = %s,
                                is_delete = 0
                            WHERE id = %s
                            """,
                            [
                                form_main_unique_id, po_num, po_date, con_contact_name, con_address, consignee_unique_id,
                                dc_received_status, dc_signed_date, ir_status, ir_signed_date, snr_status, snr_signed_date,
                                primary_product, "on" if processed_with_bg else "", "on" if processed_without_bg else "",
                                invoice_no, dc_number, invoice_date, status_code, doc_chn, reject_reason, ins_unique_id,
                                snr_verify_status, approved_by, approved_date, existing[0],
                            ],
                        )
                    else:
                        cur.execute(
                            """
                            INSERT INTO sign_doc_verification_detail
                            (
                                unique_id, form_main_unique_id, po_num, po_date, con_contact_name, con_address,
                                con_unique_id, dc_received_status, dc_signed_date, ir_status, ir_signed_date,
                                snr_status, snr_signed_date, po_product_name, with_bg, without_bg, invoice_no,
                                dc_number, invoice_date, status_app, doc_chn, reject_reason, ins_unique_id,
                                snr_verify_status, inv_verify_status, inv_verify_approvedby, inv_verify_approved_date,
                                is_active, is_delete, acc_year, session_id, sess_user_type, sess_user_id, sess_company_id, sess_branch_id
                            )
                            VALUES
                            (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, 1, %s, %s, 1, 0, '', '', '', '', '', '')
                            """,
                            [
                                unique_id, form_main_unique_id, po_num, po_date, con_contact_name, con_address,
                                consignee_unique_id, dc_received_status, dc_signed_date, ir_status, ir_signed_date,
                                snr_status, snr_signed_date, primary_product, "on" if processed_with_bg else "",
                                "on" if processed_without_bg else "", invoice_no, dc_number, invoice_date, status_code,
                                doc_chn, reject_reason, ins_unique_id, snr_verify_status, approved_by, approved_date,
                            ],
                        )

                    cur.execute(
                        """
                        UPDATE installation_details_sublist
                        SET document_verification_status = %s,
                            sign_mismatch_status = %s,
                            sign_reject_reason = %s,
                            is_delete = 0
                        WHERE po_form_unique_id = %s
                          AND consignee_unique_id = %s
                          AND invoice_no = %s
                          AND dc_number = %s
                          AND unique_id = %s
                        """,
                        [status_code, status_code if status_code == "1" else 0, reject_reason, form_main_unique_id, consignee_unique_id, invoice_no, dc_number, ins_unique_id],
                    )
                    cur.execute(
                        """
                        UPDATE installation_details
                        SET document_verification_status = %s,
                            sign_mismatch_status = %s,
                            sign_reject_reason = %s,
                            snr_verify_status = %s
                        WHERE po_form_unique_id = %s
                          AND consignee_unique_id = %s
                          AND invoice_no = %s
                          AND dc_number = %s
                          AND is_delete = 0
                        """,
                        [status_code, status_code if status_code == "1" else 0, reject_reason, 1 if status_code == "2" else snr_verify_status, form_main_unique_id, consignee_unique_id, invoice_no, dc_number],
                    )
                    cur.execute(
                        """
                        UPDATE dc_ir_doc_dispatch_details
                        SET sign_mismatch_status = %s,
                            sign_reject_reason = %s
                        WHERE po_form_unique_id = %s
                          AND consignee_unique_id = %s
                          AND invoice_no = %s
                          AND dc_number = %s
                          AND is_delete = 0
                        """,
                        [status_code if status_code == "1" else 0, reject_reason, form_main_unique_id, consignee_unique_id, invoice_no, dc_number],
                    )
                    cur.execute(
                        """
                        UPDATE invoice_creation_main
                        SET sign_mismatch_status = %s,
                            sign_reject_reason = %s
                        WHERE form_main_unique_id = %s
                          AND consignee_unique_id = %s
                          AND invoice_no = %s
                          AND dc_number = %s
                          AND is_delete = 0
                        """,
                        [status_code if status_code == "1" else 0, reject_reason, form_main_unique_id, consignee_unique_id, invoice_no, dc_number],
                    )
                    cur.execute(
                        """
                        UPDATE invoice_verfication_table
                        SET signed_complete_status = 2
                        WHERE form_main_unique_id = %s
                          AND invoice_no = %s
                        """,
                        [form_main_unique_id, invoice_no],
                    )
                    if ho_received_date:
                        cur.execute(
                            """
                            UPDATE dc_num_status
                            SET signed_complete_status = 2,
                                dc_sign_doc_date = %s
                            WHERE form_main_unique_id = %s
                              AND invoice_no = %s
                            """,
                            [ho_received_date, form_main_unique_id, invoice_no],
                        )
                    if processed_with_bg:
                        cur.execute("UPDATE po_form SET proceed_bg = 1 WHERE unique_id = %s AND is_delete = 0", [form_main_unique_id])
                    elif processed_without_bg:
                        cur.execute("UPDATE po_form SET proceed_bg = 2 WHERE unique_id = %s AND is_delete = 0", [form_main_unique_id])
        except Exception as exc:
            return Response({"status": False, "message": str(exc)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

        return Response({"status": True, "message": "Signed document verification saved successfully.", "verificationUniqueId": unique_id})


class SignDocVerificationDeleteView(APIView):
    permission_classes = [IsAuthenticated]

    def delete(self, request, unique_id):
        consignee_unique_id = str(request.query_params.get("consignee_unique_id", "") or "").strip()
        invoice_no = str(request.query_params.get("invoice_no", "") or "").strip()
        dc_number = str(request.query_params.get("dc_number", "") or "").strip()
        ins_unique_id = str(request.query_params.get("ins_unique_id", "") or "").strip()
        form_main_unique_id = str(request.query_params.get("form_main_unique_id", "") or "").strip()

        try:
            with transaction.atomic():
                with connection.cursor() as cur:
                    cur.execute(
                        """
                        SELECT form_main_unique_id, con_unique_id, invoice_no, dc_number, ins_unique_id, unique_id
                        FROM sign_doc_verification_detail
                        WHERE unique_id = %s
                          AND is_delete = 0
                        ORDER BY id DESC
                        LIMIT 1
                        """,
                        [unique_id],
                    )
                    row = cur.fetchone()
                    if not row and ins_unique_id:
                        cur.execute(
                            """
                            SELECT form_main_unique_id, con_unique_id, invoice_no, dc_number, ins_unique_id, unique_id
                            FROM sign_doc_verification_detail
                            WHERE ins_unique_id = %s
                              AND is_delete = 0
                            ORDER BY id DESC
                            LIMIT 1
                            """,
                            [ins_unique_id],
                        )
                        row = cur.fetchone()
                    if not row and form_main_unique_id and consignee_unique_id and invoice_no and dc_number:
                        cur.execute(
                            """
                            SELECT form_main_unique_id, con_unique_id, invoice_no, dc_number, ins_unique_id, unique_id
                            FROM sign_doc_verification_detail
                            WHERE form_main_unique_id = %s
                              AND con_unique_id = %s
                              AND invoice_no = %s
                              AND dc_number = %s
                              AND is_delete = 0
                            ORDER BY id DESC
                            LIMIT 1
                            """,
                            [form_main_unique_id, consignee_unique_id, invoice_no, dc_number],
                        )
                        row = cur.fetchone()
                    if not row and consignee_unique_id and invoice_no and dc_number:
                        cur.execute(
                            """
                            SELECT form_main_unique_id, con_unique_id, invoice_no, dc_number, ins_unique_id, unique_id
                            FROM sign_doc_verification_detail
                            WHERE con_unique_id = %s
                              AND invoice_no = %s
                              AND dc_number = %s
                              AND is_delete = 0
                            ORDER BY id DESC
                            LIMIT 1
                            """,
                            [consignee_unique_id, invoice_no, dc_number],
                        )
                        row = cur.fetchone()
                    bill_invoice = invoice_no
                    bill_dc = dc_number
                    if row:
                        form_main_unique_id, consignee_unique_id, invoice_no, dc_number, ins_unique_id, unique_id = row
                        bill_invoice = invoice_no
                        bill_dc = dc_number

                    if bill_invoice and bill_dc:
                        cur.execute(
                            """
                            SELECT bill_no
                            FROM sign_doc_verification_detail
                            WHERE invoice_no = %s
                              AND dc_number = %s
                              AND is_delete = 0
                            ORDER BY id DESC
                            LIMIT 1
                            """,
                            [bill_invoice, bill_dc],
                        )
                        bill_row = cur.fetchone()
                        if bill_row and str(bill_row[0] or "").strip():
                            return Response({"status": False, "message": "Cannot delete record. Bill has already been created."}, status=status.HTTP_400_BAD_REQUEST)

                    if row:
                        cur.execute("UPDATE sign_doc_verification_detail SET is_delete = 1 WHERE unique_id = %s", [unique_id])
                        cur.execute(
                            """
                            UPDATE installation_details_sublist
                            SET document_verification_status = '0',
                                sign_mismatch_status = 0,
                                sign_reject_reason = '',
                                is_delete = 0
                            WHERE po_form_unique_id = %s
                              AND consignee_unique_id = %s
                              AND invoice_no = %s
                              AND dc_number = %s
                              AND unique_id = %s
                            """,
                            [form_main_unique_id, consignee_unique_id, invoice_no, dc_number, ins_unique_id],
                        )
                        cur.execute(
                            """
                            UPDATE installation_details
                            SET document_verification_status = '0',
                                sign_mismatch_status = 0,
                                sign_reject_reason = ''
                            WHERE po_form_unique_id = %s
                              AND consignee_unique_id = %s
                              AND invoice_no = %s
                              AND dc_number = %s
                              AND is_delete = 0
                            """,
                            [form_main_unique_id, consignee_unique_id, invoice_no, dc_number],
                        )
                        cur.execute(
                            """
                            UPDATE dc_ir_doc_dispatch_details
                            SET sign_mismatch_status = 0,
                                sign_reject_reason = ''
                            WHERE po_form_unique_id = %s
                              AND consignee_unique_id = %s
                              AND invoice_no = %s
                              AND dc_number = %s
                              AND is_delete = 0
                            """,
                            [form_main_unique_id, consignee_unique_id, invoice_no, dc_number],
                        )
                        cur.execute(
                            """
                            UPDATE invoice_creation_main
                            SET sign_mismatch_status = 0,
                                sign_reject_reason = ''
                            WHERE form_main_unique_id = %s
                              AND consignee_unique_id = %s
                              AND invoice_no = %s
                              AND dc_number = %s
                              AND is_delete = 0
                            """,
                            [form_main_unique_id, consignee_unique_id, invoice_no, dc_number],
                        )
                    elif form_main_unique_id and consignee_unique_id and invoice_no and dc_number and ins_unique_id:
                        cur.execute(
                            """
                            UPDATE installation_details
                            SET dc_delivery_status = '0'
                            WHERE po_form_unique_id = %s
                              AND consignee_unique_id = %s
                              AND invoice_no = %s
                              AND dc_number = %s
                              AND is_delete = 0
                            """,
                            [form_main_unique_id, consignee_unique_id, invoice_no, dc_number],
                        )
                    else:
                        return Response({"status": False, "message": "Record not found."}, status=status.HTTP_404_NOT_FOUND)
        except Exception as exc:
            return Response({"status": False, "message": str(exc)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

        return Response({"status": True, "message": "Record deleted successfully."})


class SignDocVerificationExportView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        params = {**_parse_list_params(request), "tab": "verified", "start": 0, "length": -1}
        _, rows = _get_saved_rows(params, ["2"])
        return _build_excel_response(rows, "Signed_Document_Verified.xlsx")


class SignDocVerificationExportPendingView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        params = {**_parse_list_params(request), "tab": "mismatch", "start": 0, "length": -1}
        _, mismatch_rows = _get_saved_rows(params, ["1"])
        params_pending = {**_parse_list_params(request), "tab": "pending", "pending_type": "ir", "start": 0, "length": -1}
        _, pending_ir = _get_pending_rows(params_pending)
        params_pending["pending_type"] = "snr"
        _, pending_snr = _get_pending_rows(params_pending)
        return _build_excel_response(pending_ir + pending_snr + mismatch_rows, "Signed_Document_Pending_Mismatch.xlsx")


def _build_excel_response(rows, filename: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Signed Document Verification"
    headers = ["S.No", "PO No", "PO Date", "Followed By", "Ledger Name", "Invoice No", "Invoice Date", "DC No", "DC Date", "Status"]
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = openpyxl.styles.Font(bold=True)

    for idx, row in enumerate(rows, start=1):
        ws.append(
            [
                idx,
                row.get("poNo", ""),
                row.get("poDate", ""),
                row.get("followedBy", ""),
                row.get("ledgerName", ""),
                row.get("invoiceNo", ""),
                row.get("invoiceDate", ""),
                row.get("dcNo", ""),
                row.get("dcDate", ""),
                row.get("status", ""),
            ]
        )

    for col_cells in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col_cells), default=10)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 3, 40)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    response = HttpResponse(
        buf.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
